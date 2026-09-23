"""Convert the official NCES CIP–SOC workbook to a small CSV for the Node importer."""

import csv
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data" / "CIP2020_SOC2018_Crosswalk.xlsx"
TARGET = ROOT / "data" / "cip_soc_crosswalk.csv"

workbook = load_workbook(SOURCE, read_only=True, data_only=True)
sheet = workbook["CIP-SOC"]
with TARGET.open("w", newline="", encoding="utf-8") as target:
    writer = csv.writer(target)
    writer.writerows(sheet.values)

print(f"Prepared {sheet.max_row - 1} CIP–SOC links in {TARGET}")
